"""
Word报告生成器
生成可编辑的Word格式分析报告
"""

import os
from datetime import datetime
from typing import Dict, Any

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class WordReportGenerator:
    """Word报告生成器"""

    def __init__(self):
        """初始化Word生成器"""
        pass

    def generate(self, analysis_data: Dict[str, Any], output_path: str) -> bool:
        """
        生成Word报告

        Args:
            analysis_data: 分析结果数据
            output_path: 输出文件路径

        Returns:
            bool: 是否成功
        """
        try:
            # 创建Word文档
            doc = Document()

            # 设置中文字体
            self._set_chinese_font(doc)

            # 1. 添加标题
            self._add_title(doc)

            # 2. 添加基本信息
            self._add_basic_info(doc, analysis_data)

            # 3. 添加分析结果
            self._add_analysis_results(doc, analysis_data)

            # 4. 添加综合评价
            self._add_evaluation(doc, analysis_data)

            # 保存文档
            doc.save(output_path)

            return True

        except Exception as e:
            print(f"Word生成失败: {e}")
            return False

    def _set_chinese_font(self, doc: Document):
        """设置中文字体"""
        # 设置默认字体
        style = doc.styles['Normal']
        style.font.name = 'SimHei'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

    def _add_title(self, doc: Document):
        """添加标题"""
        # 主标题
        title = doc.add_heading('HRG谐振陀螺分析报告', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 设置标题字体
        for run in title.runs:
            run.font.name = 'SimHei'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        doc.add_paragraph()

    def _add_basic_info(self, doc: Document, data: Dict[str, Any]):
        """添加基本信息"""
        # 章节标题
        heading = doc.add_heading('一、基本信息', level=1)
        for run in heading.runs:
            run.font.name = 'SimHei'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        # 创建表格
        table = doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # 填充数据
        info_data = [
            ('项目', '内容'),
            ('报告生成时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
            ('分析软件版本', 'v3.0'),
            ('执行标准', 'GB/T 24630.1-2024, GB/T 7235-2004, GB/T 3505-2009, GB/T 6062-2009'),
        ]

        for i, (key, value) in enumerate(info_data):
            row = table.rows[i]
            row.cells[0].text = key
            row.cells[1].text = value

            # 设置字体
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'SimHei'
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        doc.add_paragraph()

    def _add_analysis_results(self, doc: Document, data: Dict[str, Any]):
        """添加分析结果"""
        # 装配误差分析
        if 'assembly_error' in data:
            self._add_assembly_error(doc, data['assembly_error'])

        # 圆度误差分析
        if 'roundness' in data:
            self._add_roundness(doc, data['roundness'])

        # 球度误差分析
        if 'sphericity' in data:
            self._add_sphericity(doc, data['sphericity'])

    def _add_assembly_error(self, doc: Document, data: Dict[str, Any]):
        """添加装配误差分析"""
        heading = doc.add_heading('二、装配误差分析', level=1)
        for run in heading.runs:
            run.font.name = 'SimHei'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        # 创建表格
        table = doc.add_table(rows=5, cols=2)
        table.style = 'Light Grid Accent 1'

        table_data = [
            ('指标', '数值'),
            ('平均距离', f"{data.get('mean_distance', 0):.6f} mm ({data.get('mean_distance', 0)*1000:.2f} μm)"),
            ('标准差', f"{data.get('std_distance', 0):.6f} mm ({data.get('std_distance', 0)*1000:.2f} μm)"),
            ('最小距离', f"{data.get('min_distance', 0):.6f} mm ({data.get('min_distance', 0)*1000:.2f} μm)"),
            ('最大距离', f"{data.get('max_distance', 0):.6f} mm ({data.get('max_distance', 0)*1000:.2f} μm)"),
        ]

        for i, (key, value) in enumerate(table_data):
            row = table.rows[i]
            row.cells[0].text = key
            row.cells[1].text = value

            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'SimHei'
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        doc.add_paragraph()

    def _add_roundness(self, doc: Document, data: Dict[str, Any]):
        """添加圆度误差分析"""
        heading = doc.add_heading('三、圆度误差分析', level=1)
        for run in heading.runs:
            run.font.name = 'SimHei'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        table = doc.add_table(rows=3, cols=2)
        table.style = 'Light Grid Accent 1'

        table_data = [
            ('指标', '数值'),
            ('最大圆度误差', f"{data.get('max_error', 0):.2f} μm"),
            ('平均圆度误差', f"{data.get('mean_error', 0):.2f} μm"),
        ]

        for i, (key, value) in enumerate(table_data):
            row = table.rows[i]
            row.cells[0].text = key
            row.cells[1].text = value

            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'SimHei'
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        doc.add_paragraph()

    def _add_sphericity(self, doc: Document, data: Dict[str, Any]):
        """添加球度误差分析"""
        heading = doc.add_heading('四、球度误差分析', level=1)
        for run in heading.runs:
            run.font.name = 'SimHei'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        table = doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'

        center = data.get('center', [0, 0, 0])
        table_data = [
            ('指标', '数值'),
            ('拟合球心', f"({center[0]:.6f}, {center[1]:.6f}, {center[2]:.6f}) mm"),
            ('拟合半径', f"{data.get('radius', 0):.6f} mm"),
            ('球度误差', f"{data.get('error', 0):.2f} μm"),
        ]

        for i, (key, value) in enumerate(table_data):
            row = table.rows[i]
            row.cells[0].text = key
            row.cells[1].text = value

            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'SimHei'
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        doc.add_paragraph()

    def _add_evaluation(self, doc: Document, data: Dict[str, Any]):
        """添加综合评价"""
        heading = doc.add_heading('五、综合评价', level=1)
        for run in heading.runs:
            run.font.name = 'SimHei'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        eval_text = """
本报告基于HRG点云分析结果生成，包含装配误差、圆度误差、球度误差等关键指标。
分析过程符合国家标准GB/T 24630.1-2024、GB/T 7235-2004、GB/T 3505-2009和GB/T 6062-2009的要求。
        """

        para = doc.add_paragraph(eval_text.strip())
        for run in para.runs:
            run.font.name = 'SimHei'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')
